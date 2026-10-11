"""Read-only inventory and text collection of all Alibaba product packages.

Source artwork is indexed, never rewritten. Collected claims keep their original
evidence and publication limits; collection does not authorize website claims.
Run with the bundled Python runtime (Pillow, pypdf and openpyxl).
"""
import csv
import hashlib
import json
import os
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

ROOT = Path("D:/work/codex/阿里国际站运营/outputs/alibaba-detail-image-studio")
OUT = Path(__file__).resolve().parents[1] / "docs/reviews/alibaba-materials-20261009"
EXCLUDE = {"node_modules", ".git", ".venv", "venv", "__pycache__"}
TEXT = {".json", ".md", ".markdown", ".txt", ".csv", ".html", ".htm", ".jsonl"}
MEDIA = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".svg", ".mp4"}
SUPPORT = {".pdf", ".xlsx", ".docx", ".zip"}
FACT_KEYS = re.compile(r"confirm|verified|fact|spec|param|evidence|crosswalk|missing|block|hold|scope|model", re.I)


def family(name):
    rules = [
        ("泡沫喷头及配件", r"泡沫|PT-series|ZSPTX"),
        ("水雾及水幕喷头", r"水雾|水幕|mist-nozzle"),
        ("干式喷头", r"干式喷头"), ("隐蔽式喷头", r"隐蔽式"),
        ("大K及ESFR喷头", r"ESFR|ZST161|ZST-K202"),
        ("扩大覆盖喷头", r"扩大覆盖"), ("快速响应喷头", r"快速响应|K-ZST"),
        ("其他洒水喷头及配件", r"喷头|洒水|sprinkler"),
        ("湿式报警阀", r"湿式报警|wet-alarm"), ("雨淋阀", r"雨淋报警阀"),
        ("预作用装置", r"预作用"), ("干式报警阀", r"干式报警阀"),
        ("水流指示器及末端试水", r"指示器|试水|测试排水"),
        ("消防蝶阀", r"蝶阀"), ("消防闸阀", r"闸阀"),
        ("软管卷盘", r"卷盘"), ("消防水枪", r"水枪"),
        ("消防水带及接口", r"接口|接头|水带|coupling"),
        ("水泵接合器及消防队进水口", r"接合器|进水口"),
        ("室外消火栓", r"室外|湿式消火栓"),
        ("室内消火栓及栓阀", r"消火栓|消防栓阀|压力限制角阀"),
    ]
    return next((label for label, pattern in rules if re.search(pattern, name, re.I)), "其他消防配套产品")


def role(path):
    n = path.name.lower()
    if any(w in n for w in ("prompt", "generation", "overlay", "quality", "audit", "qc", "自查", "质量")):
        return "制作或质量记录：不能单独作为产品事实依据"
    if any(w in n for w in ("evidence", "fact", "crosswalk", "confirm", "parameter", "核验", "确认", "事实", "参数", "报告", "report", "台账")):
        return "参数或证据记录：需逐字段检查适用范围"
    if "source" in path.parts:
        return "原始资料：需核对型号及来源"
    if "images" in path.parts or "final" in path.parts:
        return "输出图片或文案：需回溯原始证据"
    return "商品文案或辅助资料：不是独立事实证明"


def extract(path):
    if path.suffix.lower() in TEXT:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return raw.decode(encoding), {"encoding": encoding}
            except UnicodeDecodeError:
                pass
        raise ValueError("Cannot decode text without loss")
    if path.suffix.lower() == ".pdf":
        pdf = PdfReader(path)
        pages = [{"page": i + 1, "text": p.extract_text() or ""} for i, p in enumerate(pdf.pages)]
        return "\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages), {
            "pages": len(pages), "pagesWithoutExtractableText": [p["page"] for p in pages if not p["text"].strip()],
            "limit": "Text extraction only. Scanned pages need visual review/OCR; no report scope or approval verified by this collector."
        }
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            return "\n".join("[Sheet " + sheet.title + "]\n" + "\n".join(
                "\t".join("" if v is None else str(v) for v in row)
                for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)
            ) for sheet in wb), {"sheets": wb.sheetnames, "limit": "Cell contents only; formulas preserved without execution."}
        finally:
            wb.close()
    return None, {}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    packages, files, corpus, errors = [], [], [], []
    for package in sorted(p for p in ROOT.iterdir() if p.is_dir()):
        record = {"package": package.name, "family": family(package.name), "root": str(package), "facts": {}, "files": []}
        manifest = package / "manifest.json"
        if manifest.exists():
            data = json.loads(manifest.read_text(encoding="utf-8-sig"))
            record["manifest"] = str(manifest)
            record["facts"] = {k: v for k, v in data.items() if FACT_KEYS.search(k)} if isinstance(data, dict) else {"manifest": data}
        else:
            record["manifestMissing"] = True
        for current, dirs, names in os.walk(package):
            dirs[:] = sorted(d for d in dirs if d not in EXCLUDE)
            for name in sorted(names):
                path = Path(current) / name
                ext = path.suffix.lower()
                if ext not in TEXT | MEDIA | SUPPORT:
                    continue
                entry = {"package": package.name, "family": record["family"], "path": str(path),
                         "relativePath": str(path.relative_to(package)), "extension": ext,
                         "bytes": path.stat().st_size, "role": role(path), "extraction": "indexed"}
                try:
                    if ext in TEXT | {".pdf", ".xlsx"}:
                        body, details = extract(path)
                        entry["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
                        entry["extraction"] = "text-collected"
                        corpus.append({**entry, "details": details, "text": body})
                    elif ext in MEDIA - {".mp4", ".svg"}:
                        with Image.open(path) as im:
                            entry["width"], entry["height"] = im.size
                except Exception as exc:
                    entry["extraction"] = "error"
                    entry["error"] = str(exc)
                    errors.append({"path": str(path), "error": str(exc)})
                files.append(entry)
                record["files"].append(entry["path"])
        packages.append(record)
    with (OUT / "files.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["package", "family", "path", "relativePath", "extension", "bytes", "width", "height", "role", "extraction", "sha256", "error"])
        writer.writeheader()
        writer.writerows(files)
    with (OUT / "source-text.jsonl").open("w", encoding="utf-8") as handle:
        for row in corpus:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    (OUT / "packages.json").write_text(json.dumps(packages, ensure_ascii=False, indent=2), encoding="utf-8")
    stats = {"date": "2026-10-09", "root": str(ROOT), "packages": len(packages),
             "manifestMissing": [p["package"] for p in packages if p.get("manifestMissing")],
             "indexedFiles": len(files), "extensions": dict(Counter(f["extension"] for f in files)),
             "collectedTextFiles": len(corpus), "families": dict(Counter(p["family"] for p in packages)),
             "errors": errors, "pdfs": [{"path": c["path"], **c["details"]} for c in corpus if c["extension"] == ".pdf"],
             "boundary": "All product packages indexed; software dependencies and executable scripts excluded. Artwork not OCRed wholesale. Collection does not establish publication eligibility."}
    (OUT / "collection-summary.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# 阿里国际站全系列资料索引", "", "采集日期：2026-10-09。以下是内部资料汇编，不是已批准的官网文案。", "",
             "全部产品文件路径见 files.csv；可检索文案、参数表、证据记录及 PDF/XLSX 提取内容见 source-text.jsonl；清单记录见 packages.json。保留原文中的 pending/blocked 等限制。", "",
             "软件依赖和执行脚本排除；图片只建立路径和尺寸索引，未逐张完成技术审核。PDF 扫描页未作 OCR。", ""]
    for idx, package in enumerate(packages, 1):
        lines += [f"## {idx:03d} {package['package']}", "", f"系列分组：{package['family']}；收录文件：{len(package['files'])}。", "",
                  f"资料目录：[{package['package']}](<{Path(package['root']).as_posix()}>)", ""]
        if package.get("manifestMissing"):
            lines += ["缺少根目录 manifest.json；已继续收录现存图片和文案，不判定资料包为空。", ""]
        if package["facts"]:
            lines += ["原始清单中的型号、参数、证据及限制（尚需与官网逐字段核对）：", "", "```json", json.dumps(package["facts"], ensure_ascii=False, indent=2), "```", ""]
        evidence = [f for f in files if f["package"] == package["package"] and f["extraction"] == "text-collected" and "参数或证据" in f["role"]]
        if evidence:
            lines += ["参数与证据文档：", ""] + [f"- [{f['relativePath']}](<{Path(f['path']).as_posix()}>)" for f in evidence] + [""]
    (OUT / "all-products.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({k: v for k, v in stats.items() if k not in {"pdfs", "root", "boundary"}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
